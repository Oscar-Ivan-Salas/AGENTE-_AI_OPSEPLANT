import os
from docx import Document
from docx.shared import Inches
from datetime import datetime
from config import BASE_DIR, LOGO_PATH, PLANTILLA_PATH

def generar_reporte(ruta_base: str, resumen: str, nombre_supervisor: str) -> str:
    fecha = datetime.now().strftime("%Y-%m-%d")
    ruta_fotos = os.path.join(ruta_base, "fotos")
    nombre_reporte = f"Reporte_OPSEPLANT_{fecha}.docx"
    ruta_reporte = os.path.join(ruta_base, nombre_reporte)

    # Cargar plantilla
    doc = Document(PLANTILLA_PATH)

    # Insertar logo si existe
    if os.path.exists(LOGO_PATH):
        doc.sections[0].header.paragraphs[0].add_run().add_picture(LOGO_PATH, width=Inches(1.5))

    # Agregar datos básicos
    doc.add_paragraph(f"Supervisor: {nombre_supervisor}")
    doc.add_paragraph(f"Fecha: {fecha}")
    doc.add_paragraph("Resumen del avance:")
    doc.add_paragraph(resumen)

    # Insertar fotos
    if os.path.exists(ruta_fotos):
        doc.add_paragraph("Fotos del avance:")
        for nombre_foto in os.listdir(ruta_fotos):
            ruta_foto = os.path.join(ruta_fotos, nombre_foto)
            try:
                doc.add_picture(ruta_foto, width=Inches(4))
                doc.add_paragraph(nombre_foto)
            except Exception as e:
                print(f"⚠️ Error al insertar {nombre_foto}: {e}")

    # Guardar el documento
    doc.save(ruta_reporte)
    return ruta_reporte

